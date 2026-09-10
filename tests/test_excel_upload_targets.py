"""Identical upload checkpoint contract for direct and worksheet targets."""

from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock
import httpx
import pytest
from openai import OpenAI, APIConnectionError, InternalServerError
from openpyxl import Workbook
from src.ingestion.manifest_builder import build_manifest
from src.ingestion.manifest_persistence import (
    create_manifest,
    save_manifest,
    load_manifest,
    ManifestPersistenceError,
)
from src.ingestion.excel_preprocessing import preprocess_workbook
from src.ingestion.upload_targets import enumerate_upload_targets, UploadTargetKey
from src.ingestion.upload_workflow import (
    run_manifest_upload,
    prepare_manifest_upload,
    DefinitePreRemoteUploadError,
)
from src.ingestion.uploader import upload_openai_file, attach_file_and_poll
from src.ingestion.vector_store_manager import create_vector_store
from test_ingestion_recovery import fake_client, attachment


@pytest.fixture(params=["direct", "worksheet"])
def candidate(tmp_path, request):
    root = tmp_path / "VDR"
    root.mkdir()
    if request.param == "worksheet":
        book = Workbook()
        book.active.title = "Revenue"
        book.active["A1"] = "Revenue 2025 100"
        book.save(root / "model.xlsx")
        book.close()
    else:
        (root / "memo.pdf").write_bytes(b"pdf fixture")
    create_manifest(build_manifest(str(root)), root)
    if request.param == "worksheet":
        preprocess_workbook(root, "model.xlsx")
    manifest = load_manifest(root)
    manifest.vector_store_id = "vs_test"
    save_manifest(manifest, root)
    return root


def owner(root):
    return enumerate_upload_targets(load_manifest(root), root)[0].state_owner


def good_attach(_client, store, file_id):
    return attachment(file_id, store)


@pytest.mark.parametrize(
    "failure",
    [
        "none",
        "pre_remote",
        "uncertain",
        "missing_id",
        "unusable_id",
        "attach",
        "index_failed",
    ],
)
def test_remote_boundaries(candidate, failure):
    upload = Mock()
    if failure == "pre_remote":
        upload.side_effect = DefinitePreRemoteUploadError("local open")
    elif failure == "uncertain":
        upload.side_effect = TimeoutError("uncertain")
    else:
        upload.return_value = SimpleNamespace(
            id=(
                None
                if failure == "missing_id"
                else " " if failure == "unusable_id" else "file_exact"
            )
        )
    attach = Mock(
        side_effect=TimeoutError("poll") if failure == "attach" else None,
        return_value=SimpleNamespace(
            status="failed" if failure == "index_failed" else "completed",
            id="file_exact",
            vector_store_id="vs_test",
        ),
    )
    result = run_manifest_upload(
        candidate, client_factory=object, upload_file=upload, attach_file=attach
    )
    assert upload.call_count == 1
    state = owner(candidate)
    if failure == "none":
        assert result.succeeded and state.indexing_status == "completed"
    elif failure in {"pre_remote", "uncertain", "missing_id", "unusable_id"}:
        assert state.openai_file_id is None
        assert result.no_id_retryable_count == 1
        again = Mock(return_value=SimpleNamespace(id="file_retry"))
        retry = run_manifest_upload(
            candidate, client_factory=object, upload_file=again, attach_file=good_attach
        )
        assert retry.succeeded and again.call_count == 1
        assert owner(candidate).upload_attempts == 2
    else:
        assert state.openai_file_id == "file_exact"
        assert state.indexing_status == (
            "failed" if failure == "index_failed" else "in_progress"
        )
        assert result.recovery_details["openai_file_id"] == "file_exact"
        again = Mock()
        client = fake_client({"file_exact": "completed"})
        recovered = run_manifest_upload(
            candidate, client_factory=lambda: client, upload_file=again
        )
        assert recovered.succeeded
        again.assert_not_called()


@pytest.mark.parametrize("checkpoint", [1, 2, 3, 4])
def test_checkpoint_failures(candidate, checkpoint):
    calls = 0

    def saver(manifest, root):
        nonlocal calls
        calls += 1
        if calls == checkpoint:
            raise ManifestPersistenceError("injected checkpoint failure")
        return save_manifest(manifest, root)

    upload = Mock(return_value=SimpleNamespace(id="file_returned"))
    attach = Mock(side_effect=good_attach)
    result = run_manifest_upload(
        candidate,
        client_factory=object,
        upload_file=upload,
        attach_file=attach,
        manifest_saver=saver,
    )
    assert result.critically_stopped
    assert upload.call_count == (0 if checkpoint == 1 else 1)
    assert attach.call_count == (1 if checkpoint == 4 else 0)
    if checkpoint == 2:
        assert result.recovery_file_id == "file_returned"
        assert (
            result.recovery_details["artifact_id"]
            == enumerate_upload_targets(load_manifest(candidate), candidate)[
                0
            ].key.artifact_id
        )
    client = fake_client()
    run_manifest_upload(
        candidate, client_factory=lambda: client, upload_file=upload, attach_file=attach
    )
    assert upload.call_count == (2 if checkpoint == 2 else 1)


@pytest.mark.parametrize(
    "stage",
    [
        "batch_started",
        "file_started",
        "manifest_marked_uploading",
        "file_uploaded",
        "file_id_persisted",
        "attachment_started",
        "indexing_completed",
        "batch_completed",
    ],
)
def test_callback_isolation(candidate, stage):
    events = []
    observed_ids = []

    def callback(event):
        events.append(event.kind)
        if event.kind == "file_uploaded":
            observed_ids.append(owner(candidate).openai_file_id)
        if event.kind == stage:
            raise RuntimeError("cosmetic callback")

    upload = Mock(return_value=SimpleNamespace(id="file_stable"))
    result = run_manifest_upload(
        candidate,
        client_factory=object,
        upload_file=upload,
        attach_file=good_attach,
        progress_callback=callback,
    )
    assert (
        result.succeeded
        and stage in events
        and owner(candidate).openai_file_id == "file_stable"
    )
    assert upload.call_count == 1
    assert observed_ids == ["file_stable"]


@pytest.mark.parametrize("kind", ["file", "attachment", "store"])
@pytest.mark.parametrize("failure", ["timeout", "server_error"])
def test_sdk_operation_specific_retry_counts(tmp_path, kind, failure, monkeypatch):
    monkeypatch.setattr("openai._base_client.time.sleep", lambda _: None)
    requests = []

    def transport(request):
        requests.append(request)
        if failure == "timeout":
            raise httpx.ReadTimeout("synthetic", request=request)
        return httpx.Response(
            500,
            json={"error": {"message": "synthetic", "type": "server_error"}},
            request=request,
        )

    path = tmp_path / "proxy.md"
    path.write_text("fixture")
    with OpenAI(
        api_key="offline-test-key",
        base_url="https://offline.invalid/v1",
        max_retries=3,
        http_client=httpx.Client(transport=httpx.MockTransport(transport)),
    ) as client:
        with pytest.raises(Exception):
            if kind == "file":
                upload_openai_file(client, path)
            elif kind == "attachment":
                attach_file_and_poll(client, "vs_test", "file_test")
            else:
                create_vector_store(client, "Offline test")
        assert client.max_retries == 3
    assert len(requests) == (3 if kind == "file" else 1)
    assert all(request.method == "POST" for request in requests)


def test_local_open_failure_is_proven_before_sdk_call(tmp_path):
    client = Mock()
    with pytest.raises(DefinitePreRemoteUploadError):
        upload_openai_file(client, tmp_path / "missing.md")
    client.files.create.assert_not_called()


@pytest.mark.parametrize(
    "middle_outcome", ["local", "uncertain", "known_id", "integrity"]
)
def test_three_worksheet_middle_failure_and_recovery_preserve_siblings(
    tmp_path, middle_outcome
):
    root = tmp_path / "VDR"
    root.mkdir()
    book = Workbook()
    book.remove(book.active)
    for name in ["Revenue", "Customers", "Headcount"]:
        book.create_sheet(name)["A1"] = name
    book.save(root / "model.xlsx")
    book.close()
    raw_before = (root / "model.xlsx").read_bytes()
    create_manifest(build_manifest(str(root)), root)
    preprocess_workbook(root, "model.xlsx")
    m = load_manifest(root)
    m.vector_store_id = "vs_siblings"
    save_manifest(m, root)
    generated = {
        p: p.read_bytes()
        for p in (root.parent / "VDR Assistant" / "derived").rglob("*")
        if p.is_file()
    }
    middle_proxy = next(p for p in generated if p.name == "sheet_002.md")
    if middle_outcome == "integrity":
        original = generated[middle_proxy]
        middle_proxy.write_bytes(b"X" + original[1:])  # Same size, wrong SHA-256.
    upload_calls = []

    def upload(_client, path):
        upload_calls.append(path.name)
        if path.name == "sheet_002.md":
            if middle_outcome == "local":
                raise DefinitePreRemoteUploadError("not sent")
            if middle_outcome == "uncertain":
                raise TimeoutError("uncertain upload")
            return SimpleNamespace(id="file_customers")
        return SimpleNamespace(id="file_" + path.stem)

    def attach(client, store, file_id):
        if file_id == "file_customers":
            raise TimeoutError("known ID incomplete")
        return good_attach(client, store, file_id)

    result = run_manifest_upload(
        root, client_factory=object, upload_file=upload, attach_file=attach
    )
    assert result.pass_outcome == "finished" and result.completed_count == 2
    if middle_outcome == "integrity":
        assert upload_calls == ["sheet_001.md", "sheet_003.md"]
        assert (
            "integrity" in load_manifest(root).files[0].derived_artifacts[1].last_error
        )
        middle_proxy.write_bytes(generated[middle_proxy])
    children = load_manifest(root).files[0].derived_artifacts
    before = [children[i].model_dump() for i in [0, 2]]
    assert all(children[i].indexing_status == "completed" for i in [0, 2])
    plan = prepare_manifest_upload(root)
    selected = (
        plan.recovery_candidates if middle_outcome == "known_id" else plan.candidates
    )
    assert len(selected) == 1 and selected[0].key.artifact_id == children[1].artifact_id
    assert all(
        c.key.artifact_id for c in selected
    )  # Workbook parents never become candidates.
    retried = Mock(return_value=SimpleNamespace(id="file_customers"))
    client = fake_client({"file_customers": "completed"})
    recovered = run_manifest_upload(
        root,
        client_factory=lambda: client,
        upload_file=retried,
        attach_file=good_attach,
    )
    assert recovered.succeeded
    if middle_outcome == "known_id":
        retried.assert_not_called()
    else:
        assert (
            retried.call_count == 1 and retried.call_args.args[1].name == "sheet_002.md"
        )
    children = load_manifest(root).files[0].derived_artifacts
    assert [children[i].model_dump() for i in [0, 2]] == before
    assert (root / "model.xlsx").read_bytes() == raw_before
    assert all(p.read_bytes() == content for p, content in generated.items())


def test_sdk_initial_poll_deadline_never_reattaches(monkeypatch):
    from src.ingestion import uploader

    ticks = iter([0, 0, 1, 3])
    monkeypatch.setattr(
        uploader,
        "time",
        SimpleNamespace(monotonic=lambda: next(ticks), sleep=lambda _: None),
    )
    requests = []

    def transport(request):
        requests.append(request)
        return httpx.Response(
            200,
            json={
                "id": "file_test",
                "object": "vector_store.file",
                "created_at": 1,
                "vector_store_id": "vs_test",
                "status": "in_progress",
                "usage_bytes": 0,
                "last_error": None,
            },
            request=request,
        )

    with OpenAI(
        api_key="offline",
        base_url="https://offline.invalid/v1",
        http_client=httpx.Client(transport=httpx.MockTransport(transport)),
    ) as client:
        result = attach_file_and_poll(
            client, "vs_test", "file_test", max_wait_seconds=2
        )
        assert result.status == "in_progress"
    assert [r.method for r in requests] == ["POST", "GET"]


def test_transient_local_lock_after_file_create_never_reuploads(candidate, monkeypatch):
    from src.ingestion import local_io

    real_replace = local_io.os.replace
    sent = False
    blocked = False

    def replacement(source, destination):
        nonlocal blocked
        if sent and not blocked and Path(destination).name == "manifest.backup.json":
            blocked = True
            raise PermissionError("transient local sharing lock")
        return real_replace(source, destination)

    monkeypatch.setattr(local_io.os, "replace", replacement)

    def upload(*args):
        nonlocal sent
        assert not sent
        sent = True
        return SimpleNamespace(id="file_keep")

    result = run_manifest_upload(
        candidate, client_factory=object, upload_file=upload, attach_file=good_attach
    )
    assert (
        blocked and result.succeeded and owner(candidate).openai_file_id == "file_keep"
    )


@pytest.mark.parametrize("extension", [".xlsx", ".xls"])
def test_uploader_never_accepts_raw_workbooks(tmp_path, extension):
    path = tmp_path / ("raw" + extension)
    path.write_bytes(b"fixture")
    client = Mock()
    with pytest.raises(DefinitePreRemoteUploadError, match="Raw workbooks"):
        upload_openai_file(client, path)
    client.files.create.assert_not_called()
    assert path.read_bytes() == b"fixture"
