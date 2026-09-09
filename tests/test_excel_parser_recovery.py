from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
import xml.etree.ElementTree as ET
import pytest
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from src.ingestion.excel_parser import parse_capture, ExcelLimits
from src.ingestion.excel_preprocessing import preprocess_workbook
from src.ingestion.manifest_persistence import (
    create_manifest,
    load_manifest,
    save_manifest,
    ManifestPersistenceError,
)
from src.ingestion.manifest_builder import build_manifest
from test_excel_preprocessing import make_book, fingerprint


def test_cached_formula_values_errors_and_external_reference(tmp_path):
    path = tmp_path / "cached.xlsx"
    book = Workbook()
    s = book.active
    s["A1"] = "=1+1"
    s["B1"] = "=1/0"
    s["C1"] = "='[External.xlsx]Sheet1'!A1"
    s["D1"] = "=1+2"
    s["E1"] = ""
    book.save(path)
    book.close()
    # Test-only ZIP fixture construction, never a production OOXML parser.
    with ZipFile(path) as archive:
        files = {name: archive.read(name) for name in archive.namelist()}
    ns = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    tree = ET.fromstring(files["xl/worksheets/sheet1.xml"])
    for cell in tree.findall(".//m:c", ns):
        ref = cell.attrib["r"]
        if ref in {"A1", "B1", "C1"}:
            cell.find("m:v", ns).text = {"A1": "2", "B1": "#DIV/0!", "C1": "99"}[ref]
            if ref == "B1":
                cell.attrib["t"] = "e"
    files["xl/worksheets/sheet1.xml"] = ET.tostring(tree)
    with ZipFile(path, "w", ZIP_DEFLATED) as archive:
        for name, data in files.items():
            archive.writestr(name, data)
    coverage, outputs = parse_capture(path, "Finance/cached.xlsx", ExcelLimits())
    text = outputs[1].decode()
    assert "| 2 | #DIV/0! | 99 | [stored formula result unavailable] |" in text
    assert "=1+1" in text and "=1/0" in text and "External.xlsx" in text
    assert coverage[0].outcome == "included"


def test_wide_long_sparse_chart_and_grouped_hidden_columns(tmp_path):
    path = tmp_path / "shapes.xlsx"
    book = Workbook()
    s = book.active
    s.title = "Wide"
    for col in range(1, 101):
        s.cell(1, col, f"Col{col}")
        s.cell(2, col, col)
    s.column_dimensions.group("B", "D", hidden=True)
    long = book.create_sheet("Long")
    for row in range(1, 1501):
        long.cell(row, 1, f"Row{row}")
        long.cell(row, 3, row)
    sparse = book.create_sheet("Sparse")
    sparse["A1"] = "Near"
    sparse["XFD1048576"] = "Far"
    blank = book.create_sheet("EmptyString")
    blank["A1"] = ""
    chart = BarChart()
    chart.add_data(Reference(s, min_col=5, max_col=6, min_row=2, max_row=2))
    book.create_chartsheet("Chart only").add_chart(chart)
    book.save(path)
    book.close()
    coverage, outputs = parse_capture(path, "shapes.xlsx", ExcelLimits())
    assert len(outputs) == 4 and coverage[-1].reason == "Unsupported chart-only sheet"
    assert "Col2" not in outputs[1].decode().replace("Col20", "").replace(
        "Col21", ""
    ).replace("Col22", "").replace("Col23", "").replace("Col24", "").replace(
        "Col25", ""
    ).replace(
        "Col26", ""
    ).replace(
        "Col27", ""
    ).replace(
        "Col28", ""
    ).replace(
        "Col29", ""
    )
    assert outputs[1].count(b"### Rows") > 1
    assert (
        outputs[2].count(b"### Rows") > 1
        and b"| 1500 | Row1500 |  | 1500 |" in outputs[2]
    )
    assert len(outputs[3]) < 1500 and b"1048576" in outputs[3] and b"XFD" in outputs[3]


def test_parser_only_opens_capture_and_never_saves_source(tmp_path, monkeypatch):
    root = make_book(tmp_path)
    before = fingerprint(root)
    from src.ingestion import excel_parser

    original = excel_parser.load_workbook
    seen = []

    def spy(path, **kwargs):
        assert not Path(path).is_relative_to(root)
        assert Path(path).name == "source.xlsx"
        seen.append((Path(path), kwargs["data_only"]))
        return original(path, **kwargs)

    monkeypatch.setattr(excel_parser, "load_workbook", spy)
    monkeypatch.setattr(
        Workbook,
        "save",
        lambda *_: (_ for _ in ()).throw(
            AssertionError("Parser must never save source workbook")
        ),
    )
    preprocess_workbook(root, "book.xlsx")
    assert (
        len(seen) == 3
        and seen[0][0] == seen[1][0] == seen[2][0]
        and [v for _, v in seen] == [False, False, True]
    )
    assert fingerprint(root) == before


@pytest.mark.parametrize(
    "stage",
    [
        "copy",
        "hash",
        "open",
        "parse",
        "serialize",
        "proxy_write",
        "validation",
        "publication",
        "completed_save",
        "post_save_verify",
        "cleanup",
    ],
)
def test_atomic_publication_faults(tmp_path, monkeypatch, stage):
    root = make_book(tmp_path)
    before = fingerprint(root)
    from src.ingestion import excel_preprocessing as service, excel_parser

    def fail(*args, **kwargs):
        raise OSError(f"injected {stage}")

    original_open = Path.open
    original_load = service.load_manifest

    def save(manifest, folder):
        if (
            stage == "completed_save"
            and manifest.files[0].excel_preprocessing.status == "completed"
        ):
            fail()
        return save_manifest(manifest, folder)

    with monkeypatch.context() as patch:
        if stage == "copy":

            def open_path(path, *args, **kwargs):
                if path == root / "book.xlsx":
                    fail()
                return original_open(path, *args, **kwargs)

            patch.setattr(Path, "open", open_path)
        elif stage == "proxy_write":

            def open_path(path, *args, **kwargs):
                if path.suffix == ".md" and args and args[0] == "xb":
                    fail()
                return original_open(path, *args, **kwargs)

            patch.setattr(Path, "open", open_path)
        elif stage == "hash":
            patch.setattr(service, "file_sha256", fail)
        elif stage == "open":
            patch.setattr(excel_parser, "load_workbook", fail)
        elif stage == "parse":
            patch.setattr(service, "parse_capture", fail)
        elif stage == "serialize":
            patch.setattr(excel_parser, "serialize_sheet", fail)
        elif stage == "validation":
            patch.setattr(service, "validate_generation", fail)
        elif stage == "publication":
            patch.setattr(Path, "rename", fail)
        elif stage == "cleanup":
            patch.setattr(service.shutil, "rmtree", fail)
        elif stage == "post_save_verify":
            failed = False

            def load(folder):
                nonlocal failed
                m = original_load(folder)
                if m.files[0].excel_preprocessing.status == "completed" and not failed:
                    failed = True
                    fail()
                return m

            patch.setattr(service, "load_manifest", load)
        if stage in {"cleanup", "post_save_verify"}:
            assert (
                preprocess_workbook(root, "book.xlsx", manifest_saver=save)
                .files[0]
                .excel_preprocessing.status
                == "completed"
            )
        else:
            with pytest.raises(OSError):
                preprocess_workbook(root, "book.xlsx", manifest_saver=save)
            record = load_manifest(root).files[0]
            assert (
                record.excel_preprocessing.status == "failed"
                and not record.derived_artifacts
            )
    assert fingerprint(root) == before


def test_processing_checkpoint_failure_and_interrupted_retry(tmp_path):
    root = make_book(tmp_path)
    before = fingerprint(root)

    def fail(*args):
        raise ManifestPersistenceError("processing checkpoint")

    with pytest.raises(ManifestPersistenceError):
        preprocess_workbook(root, "book.xlsx", manifest_saver=fail)
    assert load_manifest(root).files[0].excel_preprocessing.status == "pending"
    assert not (root.parent / "VDR Assistant" / "derived").exists()
    m = load_manifest(root)
    m.files[0].excel_preprocessing.status = "processing"
    save_manifest(m, root)
    assert (
        preprocess_workbook(root, "book.xlsx").files[0].excel_preprocessing.status
        == "completed"
    )
    assert fingerprint(root) == before


def test_generation_collision_never_overwrites_or_adopts_on_load(tmp_path):
    root = make_book(tmp_path)
    m = preprocess_workbook(root, "book.xlsx")
    artifact = m.files[0].derived_artifacts[0]
    path = root.parent / "VDR Assistant" / artifact.proxy_relative_path
    path.write_bytes(b"corrupted generation")
    with pytest.raises(ValueError, match="collision"):
        preprocess_workbook(root, "book.xlsx")
    assert path.read_bytes() == b"corrupted generation"
    record = load_manifest(root).files[0]
    assert (
        record.excel_preprocessing.status == "failed" and not record.derived_artifacts
    )


def test_compact_huge_merge_is_rejected_before_materialization(tmp_path):
    path = tmp_path / "huge-merge.xlsx"
    book = Workbook()
    book.active["A1"] = "anchor"
    book.active.merge_cells("A1:B1")
    book.save(path)
    book.close()
    with ZipFile(path) as archive:
        files = {name: archive.read(name) for name in archive.namelist()}
    files["xl/worksheets/sheet1.xml"] = files["xl/worksheets/sheet1.xml"].replace(
        b'ref="A1:B1"', b'ref="A1:XFD1048576"'
    )
    with ZipFile(path, "w", ZIP_DEFLATED) as archive:
        for name, data in files.items():
            archive.writestr(name, data)
    with pytest.raises(ValueError, match="traversal limit"):
        parse_capture(path, "huge-merge.xlsx", ExcelLimits())
