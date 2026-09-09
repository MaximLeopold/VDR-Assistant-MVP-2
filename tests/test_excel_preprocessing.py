from datetime import date
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
import pytest
from openpyxl import Workbook
from src.ingestion.manifest_builder import build_manifest
from src.ingestion.manifest_persistence import (
    create_manifest,
    load_manifest,
    save_manifest,
)
from src.ingestion.excel_preprocessing import preprocess_workbook, exclude_workbook
from src.ingestion.excel_parser import ExcelLimits


def make_book(tmp_path):
    root = tmp_path / "VDR"
    root.mkdir()
    book = Workbook()
    s = book.active
    s.title = "Revenue"
    s.append(["Metric", "2024", "2024", 0, False, True])
    s.append(["Rate", 0.226])
    s["B2"].number_format = "0.0%"
    s["C2"] = date(2025, 1, 2)
    s["D2"] = "=D1+1"
    s["E2"] = "='[external.xlsx]Sheet1'!A1"
    s["A3"] = "HIDDEN ROW"
    s.row_dimensions[3].hidden = True
    s["G1"] = "HIDDEN COLUMN"
    s.column_dimensions["G"].hidden = True
    s.merge_cells("A5:C5")
    s["A5"] = "Merged"
    s["XFD1048576"].number_format = "0.00"
    book.create_sheet("Blank")
    book.create_sheet("Hidden")["A1"] = "SECRET"
    book["Hidden"].sheet_state = "hidden"
    book.create_sheet("Very hidden")["A1"] = "SECRET"
    book["Very hidden"].sheet_state = "veryHidden"
    book.save(root / "book.xlsx")
    book.close()
    create_manifest(build_manifest(str(root)), root)
    return root


def fingerprint(root):
    return [
        (p.relative_to(root).as_posix(), p.read_bytes(), p.stat().st_mtime_ns)
        for p in root.rglob("*")
        if p.is_file()
    ]


def test_capture_determinism_visibility_and_values(tmp_path):
    root = make_book(tmp_path)
    before = fingerprint(root)
    first = preprocess_workbook(root, "book.xlsx")
    record = first.files[0]
    assert [s.outcome for s in record.excel_preprocessing.worksheets] == [
        "included",
        "excluded",
        "excluded",
        "excluded",
    ]
    proxy = (
        root.parent / "VDR Assistant" / record.derived_artifacts[0].proxy_relative_path
    )
    data = proxy.read_bytes()
    text = data.decode()
    for expected in [
        "Workbook: book.xlsx",
        "Worksheet: Revenue",
        "22.6%",
        "2025-01-02",
        "False",
        "True",
        "| 0 |",
        "stored formula result unavailable",
        "=D1+1",
        "external.xlsx",
        "Merge A5:C5",
    ]:
        assert expected in text
    assert "SECRET" not in text and "HIDDEN" not in text
    second = preprocess_workbook(root, "book.xlsx")
    assert second.files[0] == record
    assert proxy.read_bytes() == data
    assert fingerprint(root) == before
    assert not list(
        (root.parent / "VDR Assistant" / "derived/excel/.attempts").iterdir()
    )


@pytest.mark.parametrize(
    "limits",
    [
        ExcelLimits(source_bytes=1),
        ExcelLimits(traversal_cells=1),
        ExcelLimits(proxy_bytes=10),
        ExcelLimits(total_bytes=10),
        ExcelLimits(elapsed_seconds=1e-9),
    ],
)
def test_limits_no_partial_publication(tmp_path, limits):
    root = make_book(tmp_path)
    before = fingerprint(root)
    with pytest.raises(Exception):
        preprocess_workbook(root, "book.xlsx", limits=limits)
    record = load_manifest(root).files[0]
    assert (
        record.excel_preprocessing.status == "failed" and not record.derived_artifacts
    )
    assert fingerprint(root) == before
    exclude_workbook(root, "book.xlsx", "Operator exclusion")
    assert load_manifest(root).files[0].excel_preprocessing.status == "excluded"
    preprocess_workbook(root, "book.xlsx")
    assert fingerprint(root) == before


def test_corrupt_and_frozen(tmp_path):
    root = make_book(tmp_path)
    (root / "book.xlsx").write_bytes(b"not a workbook")
    before = fingerprint(root)
    with pytest.raises(Exception):
        preprocess_workbook(root, "book.xlsx")
    assert fingerprint(root) == before
    m = load_manifest(root)
    m.vector_store_id = "vs_test"
    save_manifest(m, root)
    with pytest.raises(Exception, match="freezes"):
        preprocess_workbook(root, "book.xlsx")
    with pytest.raises(Exception, match="freezes"):
        exclude_workbook(root, "book.xlsx", "reason")


def test_completed_save_error_reloads_commit(tmp_path):
    root = make_book(tmp_path)

    def saver(m, r):
        save_manifest(m, r)
        if m.files[0].excel_preprocessing.status == "completed":
            raise OSError("after commit")

    assert (
        preprocess_workbook(root, "book.xlsx", manifest_saver=saver)
        .files[0]
        .excel_preprocessing.status
        == "completed"
    )
