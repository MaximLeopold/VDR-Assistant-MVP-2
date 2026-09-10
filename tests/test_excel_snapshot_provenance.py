import json
from pathlib import Path
from types import SimpleNamespace
import pytest
from openpyxl import Workbook
from src.ingestion.manifest_builder import build_manifest
from src.ingestion.manifest_persistence import (
    create_manifest,
    save_manifest,
    load_manifest,
    derive_manifest_paths,
    ManifestPersistenceError,
    ManifestPathError,
)
from src.ingestion.excel_preprocessing import preprocess_workbook, exclude_workbook
from src.ingestion.upload_targets import enumerate_upload_targets
from src.ingestion.upload_workflow import (
    run_manifest_upload,
    prepare_manifest_upload,
    UploadPreparationError,
)
from src.ingestion.case_readiness import assess_case_readiness
from src.ingestion.case_vector_store import (
    associate_empty_case_vector_store,
    CaseVectorStoreError,
    adopt_case_vector_store,
    ensure_case_vector_store,
)
from src.config.case_registry import (
    register_prepared_case,
    load_case_registry,
    CaseRegistryError,
)
from src.retrieval.citation_resolver import resolve_citations, build_source_references
from src.schemas.citation import Citation
from src.schemas.evidence import RetrievedSearchResult, VerifiedEvidenceExcerpt
from src.schemas.answer import VDRAnswer
from src.schemas.quotation import VerifiedQuote
from src.schemas.evidence_presentation import (
    VerifiedEvidencePresentation,
    VerifiedEvidenceMetric,
)
from src.ui import chat
from test_chat_ui import FakeStreamlit


def case(tmp_path, kind="mixed"):
    root = tmp_path / "Case" / "VDR"
    root.mkdir(parents=True)
    if kind != "excel":
        (root / "memo.pdf").write_bytes(b"fixture")
    if kind != "direct":
        for folder in ["Finance", "People"]:
            (root / folder).mkdir()
            book = Workbook()
            book.active.title = "Revenue"
            book.active["A1"] = "Revenue 2025 100"
            book.create_sheet("Headcount")["A1"] = "Headcount 20"
            book.save(root / folder / "model.xlsx")
            book.close()
    create_manifest(build_manifest(str(root)), root)
    for f in load_manifest(root).files:
        if f.classification_status == "preprocess":
            preprocess_workbook(root, f.relative_path)
    m = load_manifest(root)
    m.vector_store_id = "vs_fixture"
    save_manifest(m, root)
    n = iter(range(100))
    result = run_manifest_upload(
        root,
        client_factory=object,
        upload_file=lambda *_: SimpleNamespace(id=f"file_{next(n)}"),
        attach_file=lambda _client, store, file_id: SimpleNamespace(
            status="completed", id=file_id, vector_store_id=store
        ),
    )
    assert result.succeeded, (result.message, result.recovery_details)
    return root


@pytest.mark.parametrize("kind", ["direct", "excel", "mixed"])
def test_readiness_sealing_registration_and_no_local_source_dependency(
    tmp_path, kind, monkeypatch
):
    root = case(tmp_path, kind)
    assert assess_case_readiness(root).is_ready
    registry = tmp_path / "cases.json"
    registry.write_text(
        json.dumps({"cases": [{"case_id": "unavailable", "vdr_folder": "unavailable"}]})
    )

    # A proxy may disappear after indexing; readiness and opening must not read it.
    def forbidden(*args, **kwargs):
        raise AssertionError("Source/proxy access during registration or case opening")

    monkeypatch.setattr("src.ingestion.excel_preprocessing.file_sha256", forbidden)
    monkeypatch.setattr("src.ingestion.upload_targets.file_sha256", forbidden)
    for target in enumerate_upload_targets(load_manifest(root), root):
        target.local_path.unlink()
    result = register_prepared_case(registry, "fixture", root)
    assert (
        result.status == "registered" and load_manifest(root).snapshot_state == "sealed"
    )
    assert load_case_registry(registry)[1].is_ready
    assert (
        register_prepared_case(registry, "fixture", root).status == "already_registered"
    )
    with pytest.raises(ManifestPersistenceError, match="Sealed"):
        save_manifest(load_manifest(root), root)
    with pytest.raises(UploadPreparationError, match="Sealed"):
        prepare_manifest_upload(root)


def test_registration_failure_keeps_seal_and_retries(tmp_path, monkeypatch):
    root = case(tmp_path, "excel")
    registry = tmp_path / "cases.json"
    registry.write_text(
        json.dumps({"cases": [{"case_id": "old", "vdr_folder": "old"}]})
    )
    from src.config import case_registry

    original = case_registry._atomic_replace_registry
    monkeypatch.setattr(
        case_registry,
        "_atomic_replace_registry",
        lambda *_: (_ for _ in ()).throw(
            CaseRegistryError("injected registry failure")
        ),
    )
    with pytest.raises(CaseRegistryError):
        register_prepared_case(registry, "fixture", root)
    assert load_manifest(root).snapshot_state == "sealed"
    monkeypatch.setattr(case_registry, "_atomic_replace_registry", original)
    assert register_prepared_case(registry, "fixture", root).status == "registered"


@pytest.mark.parametrize("status", ["pending", "processing", "failed", "excluded"])
def test_preprocessing_readiness_blockers(tmp_path, status):
    root = tmp_path / "VDR"
    root.mkdir()
    (root / "book.xlsx").write_bytes(b"fixture")
    m = build_manifest(str(root))
    m.vector_store_id = "vs_test"
    m.files[0].excel_preprocessing.status = status
    if status == "excluded":
        m.files[0].excel_preprocessing.exclusion_reason = "explicit"
    create_manifest(m, root)
    readiness = assess_case_readiness(root)
    assert not readiness.is_ready
    assert any(
        "Excel preprocessing" in reason for reason in readiness.blocking_reasons
    ) == (status != "excluded")
    assert any(
        "At least one searchable target" in reason
        for reason in readiness.blocking_reasons
    )


def test_frozen_generation_and_association_cannot_change(tmp_path):
    root = case(tmp_path, "excel")
    for action in ["generation", "coverage", "inventory", "association", "exclusion"]:
        m = load_manifest(root)
        if action == "generation":
            m.files[0].excel_preprocessing.transformation_version = "changed"
        elif action == "coverage":
            m.files[0].excel_preprocessing.worksheets[0].reason = "changed"
        elif action == "inventory":
            m.files[0].size_bytes += 1
        elif action == "association":
            m.vector_store_id = None
        else:
            m.files[0].excel_preprocessing.last_error = "changed exclusion state"
        with pytest.raises(ManifestPersistenceError, match="freezes"):
            save_manifest(m, root)
    with pytest.raises(ManifestPersistenceError, match="freezes"):
        exclude_workbook(root, "Finance/model.xlsx", "reason")


@pytest.mark.parametrize(
    "root_kind", ["equal", "managed_inside_raw", "raw_inside_managed"]
)
def test_disjoint_roots(root_kind, tmp_path):
    from src.ingestion.paths import validate_disjoint_roots

    raw = tmp_path / "raw"
    managed = (
        raw
        if root_kind == "equal"
        else raw / "managed" if root_kind == "managed_inside_raw" else tmp_path
    )
    with pytest.raises(ValueError, match="disjoint"):
        validate_disjoint_roots(raw, managed)


def test_assistant_root_as_raw_is_rejected(tmp_path):
    root = tmp_path / "VDR Assistant"
    root.mkdir()
    with pytest.raises(ManifestPathError):
        derive_manifest_paths(root)


@pytest.mark.parametrize("kind", ["direct", "excel", "mixed"])
def test_exact_provenance_evidence_and_replay(tmp_path, kind, monkeypatch):
    root = case(tmp_path, kind)
    m = load_manifest(root)
    targets = enumerate_upload_targets(m, root)
    citations = [
        Citation(file_id=t.state_owner.openai_file_id, filename="sheet_001.md")
        for t in targets
    ]
    resolved = resolve_citations(citations, m)
    for target, source in zip(targets, resolved):
        assert source.file_id == target.state_owner.openai_file_id
        assert "sheet_001.md" not in source.display_name
        if target.artifact:
            assert (
                source.excel_provenance.original_relative_path
                == target.source.relative_path
            )
            assert (
                source.excel_provenance.worksheet_name == target.artifact.worksheet_name
            )
            assert source.display_name.endswith(" → " + target.artifact.worksheet_name)
        else:
            assert source.excel_provenance is None
    raw = "  Revenue\t100\r\n\r\n2025  "
    results = [
        RetrievedSearchResult(file_id=c.file_id, text=raw, score=0.8) for c in citations
    ]
    results.append(
        RetrievedSearchResult(file_id="uncited", text="must not appear", score=1)
    )
    sources = build_source_references(citations, resolved, results)
    assert all(a is b for a, b in zip(sources, resolved))
    for source in sources:
        source.evidence_selection_status = "completed"
        source.selected_evidence = [
            VerifiedEvidenceExcerpt(
                role="best_support", passage_index=0, text="Revenue 100"
            )
        ]
        source.presentations = [
            VerifiedEvidencePresentation(
                passage_index=0,
                metrics=[
                    VerifiedEvidenceMetric(
                        label="Revenue", value="100", source_text="Revenue 100"
                    )
                ],
            )
        ]
    quotes = [
        VerifiedQuote(
            file_id=s.file_id, source_display_name=s.display_name, text="Revenue 100"
        )
        for s in sources[:3]
    ]
    answer = VDRAnswer(
        answer="Revenue is 100.",
        source_files=[s.display_name for s in sources],
        sources=sources,
        verified_quotes=quotes,
    )
    message = json.loads(json.dumps(chat.build_assistant_message(answer)))
    assert VDRAnswer.model_validate(message["vdr_answer"]) == answer

    def forbidden(*args, **kwargs):
        raise AssertionError("Replay must only use its stored payload.")

    monkeypatch.setattr("src.ingestion.manifest_persistence.load_manifest", forbidden)
    monkeypatch.setattr("src.retrieval.citation_resolver.resolve_citations", forbidden)
    monkeypatch.setattr(
        "src.retrieval.openai_file_search.search_vector_store", forbidden
    )
    monkeypatch.setattr("src.retrieval.openai_file_search.get_openai_client", forbidden)
    monkeypatch.setattr("openpyxl.load_workbook", forbidden)
    monkeypatch.setattr(Path, "open", forbidden)
    ui = FakeStreamlit()
    monkeypatch.setattr(chat, "st", ui)
    chat.render_chat_history([message])
    assert [text for text, options in ui.code_calls] == [raw] * len(sources)
    assert len(ui.text_calls) >= len(quotes)
    if any(
        q.file_id in {s.file_id for s in sources if s.excel_provenance} for q in quotes
    ):
        assert "Excel-derived excerpts" in ui.expander_labels
        assert "Source type: Excel-derived search representation" in ui.caption_calls


def test_unknown_and_ambiguous_ids_never_trust_proxy_filename(tmp_path):
    root = case(tmp_path, "excel")
    m = load_manifest(root)
    unknown = resolve_citations(
        [Citation(file_id="unknown", filename="sheet_001.md")], m
    )[0]
    assert unknown.display_name == "Unknown source" and unknown.excel_provenance is None
    first = m.files[0].derived_artifacts[0]
    m.files[1].derived_artifacts[0].openai_file_id = first.openai_file_id
    ambiguous = resolve_citations(
        [Citation(file_id=first.openai_file_id, filename="model.xlsx")], m
    )[0]
    assert (
        ambiguous.display_name == "Unknown source"
        and ambiguous.excel_provenance is None
    )
