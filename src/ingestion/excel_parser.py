"""Deterministic extraction from an assistant-managed captured workbook."""

from dataclasses import dataclass
from datetime import date, datetime, time, timedelta
from pathlib import Path
import re
import time as clock
from zipfile import ZipFile

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.worksheet._read_only import ReadOnlyWorksheet
from openpyxl.worksheet._reader import WorkSheetParser
from src.ingestion.manifest import WorksheetCoverage

TRANSFORMATION_VERSION = "excel_proxy_v1"


@dataclass(frozen=True)
class ExcelLimits:
    source_bytes: int = 50 * 1024 * 1024
    expanded_bytes: int = 256 * 1024 * 1024
    traversal_cells: int = 1_000_000
    worksheets: int = 1000
    proxy_bytes: int = 16 * 1024 * 1024
    total_bytes: int = 64 * 1024 * 1024
    elapsed_seconds: float = 120
    section_bytes: int = 8000
    column_window: int = 12

    def __post_init__(self):
        if any(value <= 0 for value in vars(self).values()):
            raise ValueError("Excel limits must be positive.")


class ExcelLimitError(ValueError):
    pass


class Budget:
    def __init__(self, limits):
        self.limits, self.started, self.cells = limits, clock.monotonic(), 0

    def check(self, cells=0):
        self.cells += cells
        if self.cells > self.limits.traversal_cells:
            raise ExcelLimitError(
                "Excel cell traversal limit exceeded; reduce the workbook or raise traversal_cells."
            )
        if clock.monotonic() - self.started > self.limits.elapsed_seconds:
            raise ExcelLimitError(
                "Excel preprocessing time limit exceeded; reduce the workbook or raise elapsed_seconds."
            )


def escape(value):
    return (
        str(value)
        .replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
        .replace("\\", "\\\\")
        .replace("|", "\\|")
        .replace("\r\n", "<br>")
        .replace("\r", "<br>")
        .replace("\n", "<br>")
    )


def stored_value(value, fmt):
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, timedelta):
        return str(value)
    if isinstance(value, (int, float)):
        if re.fullmatch(r"0(?:\.0+)?%", fmt):
            digits = len(fmt.split(".")[1].rstrip("%")) if "." in fmt else 0
            return f"{value * 100:.{digits}f}%"
        if re.fullmatch(r"(?:#,##)?0(?:\.0+)?", fmt):
            digits = len(fmt.split(".")[1]) if "." in fmt else 0
            return format(value, f'{"," if "," in fmt else ""}.{digits}f')
        return str(value)
    return str(value)


def serialize_sheet(sheet, cached, source_path, limits, budget):
    hidden_columns = [
        (d.min, d.max) for d in sheet.column_dimensions.values() if d.hidden
    ]

    def visible(row, col):
        return not (
            row in sheet.row_dimensions and sheet.row_dimensions[row].hidden
        ) and not any(lo <= col <= hi for lo, hi in hidden_columns)

    # openpyxl 3.1.5 materializes populated/style cells here. Unlike iter_rows,
    # this does not expand a sparse nominal dimension into millions of blanks.
    cells = {}
    for (row, col), cell in sheet._cells.items():
        budget.check(1)
        if visible(row, col) and (
            cell.value is not None or cell.data_type == "inlineStr"
        ):
            cells[row, col] = cell
    if not cells:
        return None
    parts = [
        f"# Workbook: {escape(Path(source_path).name)}\n## Source path: {escape(source_path)}\n## Worksheet: {escape(sheet.title)}\n\nDeterministic search representation. Formulas are not recalculated.\n"
    ]
    size = len(parts[0].encode("utf-8"))

    def append(text):
        nonlocal size
        size += len(text.encode("utf-8"))
        if size > limits.proxy_bytes:
            raise ExcelLimitError(
                "Worksheet proxy byte limit exceeded; reduce the workbook or raise proxy_bytes."
            )
        parts.append(text)

    # Only occupied column windows are visited. Original column letters and
    # row numbers make omitted all-blank areas explicit, without collapsing cells.
    windows = {}
    for (row, col), cell in cells.items():
        budget.check(1)
        windows.setdefault((col - 1) // limits.column_window, {})[row, col] = cell
    for window in sorted(windows):
        selected = windows[window]
        lo, hi = min(c for r, c in selected), max(c for r, c in selected)
        columns = [
            c
            for c in range(lo, hi + 1)
            if not any(a <= c <= b for a, b in hidden_columns)
        ]
        rows = sorted({r for r, c in selected})
        # Bound structural blank expansion between populated rows.
        expanded = []
        for row in rows:
            if expanded and row - expanded[-1] <= 50:
                budget.check(row - expanded[-1] - 1)
                expanded.extend(
                    r for r in range(expanded[-1] + 1, row) if visible(r, lo)
                )
            expanded.append(row)
        section, metadata, section_size, section_start = [], [], 0, None

        def flush(end):
            if not section:
                return
            heading = f"\n### Rows {section_start}-{end} | Columns {get_column_letter(lo)}-{get_column_letter(hi)}\n\n"
            header = (
                "| Excel row | "
                + " | ".join(get_column_letter(c) for c in columns)
                + " |\n"
            )
            rule = "|---|" + "---|" * len(columns) + "\n"
            append(heading + header + rule + "".join(section))
            if metadata:
                append("\n" + "\n".join(metadata) + "\n")

        last = None
        for row in expanded:
            budget.check(len(columns))
            values, notes = [], []
            for col in columns:
                cell = selected.get((row, col))
                if cell is None:
                    values.append("")
                    continue
                value = cell.value
                if cell.data_type == "f":
                    value = cached.cell(row, col).value
                    rendered = (
                        "[stored formula result unavailable]"
                        if value is None
                        else stored_value(value, cell.number_format)
                    )
                    formula = (
                        cell.value
                        if isinstance(cell.value, str)
                        else getattr(
                            cell.value, "text", "[formula metadata unavailable]"
                        )
                    )
                    notes.append(f"Formula {cell.coordinate}: {escape(formula)}")
                else:
                    rendered = stored_value(value, cell.number_format)
                values.append(escape(rendered))
                if cell.number_format != "General":
                    notes.append(
                        f"Format {cell.coordinate}: {escape(cell.number_format)}; stored value: {escape(value)}"
                    )
            line = f"| {row} | " + " | ".join(values) + " |\n"
            volume = len((line + "\n".join(notes)).encode("utf-8"))
            if section and (
                section_size + volume > limits.section_bytes or row != last + 1
            ):
                flush(last)
                section = []
                metadata = []
                section_size = 0
                section_start = None
            if section_start is None:
                section_start = row
            section.append(line)
            metadata.extend(notes)
            section_size += volume
            last = row
        flush(last)
    for merged in sorted(
        sheet.merged_cells.ranges, key=lambda m: (m.min_row, m.min_col)
    ):
        budget.check()
        if (merged.min_row, merged.min_col) in cells:
            append(
                f"\nMerge {merged}: anchor {get_column_letter(merged.min_col)}{merged.min_row}.\n"
            )
    return "".join(parts).encode("utf-8")


def guard_materialization(capture, budget):
    """Use openpyxl's streaming parser before allocating editable views.

    These pinned 3.1.5 engine interfaces expose actual stored rows and merge /
    hyperlink ranges without expanding the nominal worksheet rectangle. No
    application OOXML parsing is performed. The two extraction views remain
    formula/structure and cached/stored-result views of this same capture.
    """
    workbook = load_workbook(capture, read_only=True, data_only=False, keep_links=False)
    try:
        if len(workbook.sheetnames) > budget.limits.worksheets:
            raise ExcelLimitError("Worksheet traversal limit exceeded.")
        for worksheet in workbook:
            budget.check()
            if not isinstance(worksheet, ReadOnlyWorksheet):
                continue
            with worksheet._get_source() as stream:
                parser = WorkSheetParser(
                    stream,
                    worksheet._shared_strings,
                    data_only=False,
                    epoch=workbook.epoch,
                    date_formats=workbook._date_formats,
                    timedelta_formats=workbook._timedelta_formats,
                )
                for _, row in parser.parse():
                    budget.check(2 * len(row) + 1)
                ranges = []
                if parser.merged_cells:
                    ranges.extend(item.ref for item in parser.merged_cells.mergeCell)
                ranges.extend(item.ref for item in parser.hyperlinks.hyperlink)
                for reference in ranges:
                    lo_col, lo_row, hi_col, hi_row = range_boundaries(reference)
                    if None in (lo_col, lo_row, hi_col, hi_row):
                        raise ExcelLimitError(
                            "Unbounded worksheet range cannot be safely materialized."
                        )
                    budget.check(2 * (hi_col - lo_col + 1) * (hi_row - lo_row + 1))
    finally:
        workbook.close()


def parse_capture(capture: Path, source_path: str, limits: ExcelLimits):
    budget = Budget(limits)
    if capture.stat().st_size > limits.source_bytes:
        raise ExcelLimitError(
            "Source workbook byte limit exceeded; raise source_bytes or reduce the workbook."
        )
    with ZipFile(capture) as archive:
        if sum(i.file_size for i in archive.infolist()) > limits.expanded_bytes:
            raise ExcelLimitError(
                "Expanded workbook byte limit exceeded; reduce the workbook."
            )
    guard_materialization(capture, budget)
    formula = cached = None
    try:
        formula = load_workbook(capture, data_only=False, keep_links=False)
        budget.check()
        cached = load_workbook(capture, data_only=True, keep_links=False)
        coverage, outputs, total = [], {}, 0
        if len(formula.sheetnames) > limits.worksheets:
            raise ExcelLimitError("Worksheet traversal limit exceeded.")
        for index, name in enumerate(formula.sheetnames, 1):
            budget.check()
            sheet = formula[name]
            reason = None
            if sheet.sheet_state != "visible":
                reason = (
                    "Very hidden worksheet"
                    if sheet.sheet_state == "veryHidden"
                    else "Hidden worksheet"
                )
            elif not isinstance(sheet, Worksheet):
                reason = "Unsupported chart-only sheet"
            if reason is None:
                data = serialize_sheet(sheet, cached[name], source_path, limits, budget)
                if data is None:
                    reason = "Blank worksheet"
                else:
                    total += len(data)
                    if total > limits.total_bytes:
                        raise ExcelLimitError(
                            "Total generated workbook byte limit exceeded; reduce the workbook or raise total_bytes."
                        )
                    outputs[index] = data
            coverage.append(
                WorksheetCoverage(
                    worksheet_name=name,
                    worksheet_index=index,
                    outcome="excluded" if reason else "included",
                    reason=reason or "Visible worksheet with meaningful content",
                )
            )
        return coverage, outputs
    finally:
        if cached:
            cached.close()
        if formula:
            formula.close()
